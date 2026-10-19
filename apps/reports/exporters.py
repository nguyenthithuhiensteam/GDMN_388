"""Xuất dữ liệu ra Excel (openpyxl), PDF (reportlab, font DejaVu hỗ trợ tiếng Việt) và Word (python-docx)."""
import datetime as dt
import io
from pathlib import Path
from urllib.parse import quote

from django.http import HttpResponse
from django.utils import timezone

FONT_DIR = Path(__file__).resolve().parent / "fonts"
_fonts_ready = False


def _filename(base, ext):
    stamp = timezone.localtime().strftime("%Y%m%d-%H%M")
    return f"{base}-{stamp}.{ext}"


def _attach(resp, name):
    resp["Content-Disposition"] = f"attachment; filename*=UTF-8''{quote(name)}"
    resp["Cache-Control"] = "private, no-store"
    return resp


def _fmt(v):
    if isinstance(v, dt.datetime):
        return timezone.localtime(v).strftime("%d/%m/%Y %H:%M") if timezone.is_aware(v) else v.strftime("%d/%m/%Y %H:%M")
    if isinstance(v, dt.date):
        return v.strftime("%d/%m/%Y")
    if v is None:
        return ""
    return v


def excel_response(base_name, title, headers, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Xuất lúc: {timezone.localtime():%d/%m/%Y %H:%M}"])
    ws.append([])
    ws.append(headers)
    for cell in ws[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="3A9A3E")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows:
        ws.append([_safe_cell(_fmt(v)) for v in row])
    for i, _ in enumerate(headers, start=1):
        ws.column_dimensions[ws.cell(row=4, column=i).column_letter].width = 22
    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(),
                        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    return _attach(resp, _filename(base_name, "xlsx"))


def _safe_cell(v):
    """Chống chèn công thức (CSV/Excel injection)."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        return "'" + v
    return v


def _register_fonts():
    global _fonts_ready
    if _fonts_ready:
        return
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    pdfmetrics.registerFont(TTFont("DejaVu", str(FONT_DIR / "DejaVuSans.ttf")))
    pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(FONT_DIR / "DejaVuSans-Bold.ttf")))
    _fonts_ready = True


def _pdf_styles():
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet

    _register_fonts()
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("t", parent=base["Title"], fontName="DejaVu-Bold", fontSize=14, leading=18),
        "h": ParagraphStyle("h", parent=base["Heading3"], fontName="DejaVu-Bold", fontSize=11),
        "body": ParagraphStyle("b", parent=base["BodyText"], fontName="DejaVu", fontSize=10, leading=14),
        "small": ParagraphStyle("s", parent=base["BodyText"], fontName="DejaVu", fontSize=8, leading=10),
        "center": ParagraphStyle("c", parent=base["BodyText"], fontName="DejaVu-Bold", fontSize=10, alignment=1),
    }


def _esc(text):
    from xml.sax.saxutils import escape

    return escape(str(text)).replace("\n", "<br/>")


def pdf_table_response(base_name, title, sections):
    """sections: list of (tiêu đề, headers, rows)."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    st = _pdf_styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=30, rightMargin=30, topMargin=30, bottomMargin=30,
                            title=title)
    story = [Paragraph(_esc(title), st["title"]),
             Paragraph(f"Xuất lúc: {timezone.localtime():%d/%m/%Y %H:%M}", st["small"]), Spacer(1, 10)]
    for sec_title, headers, rows in sections:
        story.append(Paragraph(_esc(sec_title), st["h"]))
        data = [[Paragraph(_esc(h), st["small"]) for h in headers]]
        data += [[Paragraph(_esc(_fmt(v)), st["small"]) for v in r] for r in rows]
        t = Table(data, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8F5E9")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story += [t, Spacer(1, 12)]
    doc.build(story)
    return _attach(HttpResponse(buf.getvalue(), content_type="application/pdf"), _filename(base_name, "pdf"))


def _template_parts(plan, template):
    from apps.core.models import SiteSettings

    site = SiteSettings.load()
    left = (template.header_left if template and template.header_left else
            "\n".join(x for x in [site.governing_body.upper(), site.school_name.upper()] if x))
    right = template.header_right if template else "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập - Tự do - Hạnh phúc"
    prefix = template.title_prefix if template else "KẾ HOẠCH"
    sign = template.signature_title if template else ""
    footer = template.footer_note if template else ""
    status_note = {
        "draft": "DỰ THẢO", "consulting": "DỰ THẢO - LẤY GÓP Ý", "pending": "DỰ THẢO - CHỜ DUYỆT",
        "approved": "ĐÃ DUYỆT (CHƯA BAN HÀNH)", "official": "", "archived": "LƯU TRỮ",
    }[plan.status]
    return left, right, prefix, sign, footer, status_note


def plan_docx_response(plan, template=None):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    left, right, prefix, sign, footer, status_note = _template_parts(plan, template)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(13)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = left
    table.cell(0, 1).text = right
    for cell in table.rows[0].cells:
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
    if status_note:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run(status_note)
        r.bold = True
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = h.add_run(f"{prefix}\n{plan.title.upper()}" if prefix else plan.title.upper())
    run.bold = True
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"Năm học {plan.school_year}" + (f" - {plan.school_class}" if plan.school_class else "")).italic = True
    for para in (plan.content or "").splitlines():
        doc.add_paragraph(para)
    if sign:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.add_run(sign).bold = True
    if footer:
        doc.add_paragraph(footer)
    buf = io.BytesIO()
    doc.save(buf)
    resp = HttpResponse(buf.getvalue(),
                        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    return _attach(resp, _filename(f"ke-hoach-{plan.pk}", "docx"))


def plan_pdf_response(plan, template=None):
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table

    st = _pdf_styles()
    left, right, prefix, sign, footer, status_note = _template_parts(plan, template)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=50, rightMargin=40, topMargin=40, bottomMargin=40,
                            title=plan.title)
    story = [Table([[Paragraph(_esc(left), st["center"]), Paragraph(_esc(right), st["center"])]], colWidths=[240, 260])]
    if status_note:
        story.append(Paragraph(f"<para alignment='right'><b>{_esc(status_note)}</b></para>", st["body"]))
    story += [Spacer(1, 12), Paragraph(_esc(f"{prefix} {plan.title}".strip().upper()), st["title"]),
              Paragraph(_esc(f"Năm học {plan.school_year}"), st["center"]), Spacer(1, 12)]
    for para in (plan.content or "").split("\n\n"):
        story.append(Paragraph(_esc(para), st["body"]))
        story.append(Spacer(1, 6))
    if sign:
        story += [Spacer(1, 20), Paragraph(f"<para alignment='right'><b>{_esc(sign)}</b></para>", st["body"])]
    if footer:
        story += [Spacer(1, 20), Paragraph(_esc(footer), st["small"])]
    doc.build(story)
    return _attach(HttpResponse(buf.getvalue(), content_type="application/pdf"), _filename(f"ke-hoach-{plan.pk}", "pdf"))
